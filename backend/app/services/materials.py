"""Генерация материалов урока по учебнику (Модуль 5.2).

По теме и найденным фрагментам учебника (RAG) за один запрос к модели:
план урока на 45 минут, Stories-урок (8–12 слайдов), контрольная в 2–4 вариантах
и ключ ответов. Страницы-источники — только из реально найденных фрагментов.
Учитель редактирует всё в интерфейсе и выгружает в DOCX (PDF — печатью страницы).
"""
from __future__ import annotations

import asyncio
import io
import logging
from typing import Protocol

from sqlalchemy import desc, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings, get_settings
from app.db.models import LessonMaterial, Textbook, User
from app.services.gemini import SUBJECTS, GeminiError, extract_json, language_rule
from app.services.rag import Embedder, Fragment, RagError, make_embedder, retrieve

log = logging.getLogger(__name__)

LESSON_MINUTES = 45


class MaterialError(Exception):
    def __init__(self, code: str, status: int = 400):
        super().__init__(code)
        self.code = code
        self.status = status


# ---------- Проверка ответа модели ----------

def _s(value, limit: int = 2000) -> str | None:
    return value.strip()[:limit] if isinstance(value, str) and value.strip() else None


def validate_material(data, variants: int, allowed_pages: set[int] | None = None) -> dict | None:
    """Строгая схема. allowed_pages=None — правка учителем (источники не пересчитываем)."""
    if not isinstance(data, dict):
        return None
    plan, stories, test = data.get("plan"), data.get("stories"), data.get("test")
    if not isinstance(plan, dict) or not isinstance(stories, list) or not isinstance(test, dict):
        return None

    stages = []
    for st in plan.get("stages") or []:
        if not isinstance(st, dict) or not _s(st.get("name")) or not _s(st.get("activity")):
            return None
        minutes = st.get("minutes")
        if not isinstance(minutes, int) or isinstance(minutes, bool) or minutes <= 0:
            return None
        stages.append({"name": _s(st["name"], 200), "minutes": minutes, "activity": _s(st["activity"])})
    if not 3 <= len(stages) <= 8:
        return None
    total = sum(s["minutes"] for s in stages)
    if not 38 <= total <= 52:
        return None
    stages[-1]["minutes"] += LESSON_MINUTES - total  # ровно 45 минут: подгоняем последний этап
    if stages[-1]["minutes"] <= 0:
        return None

    slides = []
    for sl in stories:
        if not isinstance(sl, dict) or not _s(sl.get("title")) or not _s(sl.get("text")):
            return None
        slide = {"title": _s(sl["title"], 200), "text": _s(sl["text"], 600), "illustration": _s(sl.get("illustration"), 200) or ""}
        options, correct = sl.get("options"), sl.get("correct")
        if _s(sl.get("question")) and isinstance(options, list) and 2 <= len(options) <= 4 and all(_s(o) for o in options) \
                and isinstance(correct, int) and not isinstance(correct, bool) and 0 <= correct < len(options):
            slide.update(question=_s(sl["question"], 300), options=[_s(o, 120) for o in options], correct=correct)
        slides.append(slide)
    if not 8 <= len(slides) <= 12:
        return None

    out_variants = []
    for var in test.get("variants") or []:
        if not isinstance(var, dict):
            return None
        tasks = []
        for task in var.get("tasks") or []:
            if not isinstance(task, dict) or not _s(task.get("text")) or not _s(task.get("answer")):
                return None
            points = task.get("points", 1)
            points = points if isinstance(points, int) and not isinstance(points, bool) and 1 <= points <= 20 else 1
            tasks.append({"text": _s(task["text"], 800), "points": points, "answer": _s(task["answer"], 400)})
        if not 3 <= len(tasks) <= 10:
            return None
        out_variants.append({"name": _s(var.get("name"), 20) or chr(ord("A") + len(out_variants)), "tasks": tasks})
    if len(out_variants) != variants:
        return None

    goals = [g for g in (_s(x, 300) for x in plan.get("goals") or []) if g][:5]
    raw_sources = data.get("sources") or []
    sources = sorted({p for p in raw_sources if isinstance(p, int) and not isinstance(p, bool) and p > 0})
    if allowed_pages is not None:
        # Не даём модели «придумать» страницу: только те, что были в найденных фрагментах
        sources = [p for p in sources if p in allowed_pages] or sorted(allowed_pages)[:5]
    return {
        "plan": {"title": _s(plan.get("title"), 200) or "", "goals": goals, "stages": stages, "homework": _s(plan.get("homework"), 600) or ""},
        "stories": slides,
        "test": {"variants": out_variants},
        "answer_key": [
            {"variant": v["name"], "n": i + 1, "answer": t["answer"], "points": t["points"]}
            for v in out_variants
            for i, t in enumerate(v["tasks"])
        ],
        "sources": sources,
    }


# ---------- Генераторы ----------

class MaterialGenerator(Protocol):
    async def generate(self, topic: str, fragments: list[Fragment], variants: int, book: Textbook, lang: str) -> dict: ...


SYSTEM_PROMPT = (
    "Ты — методист и помощник учителя школы Узбекистана. Готовишь материалы урока СТРОГО по "
    "фрагментам учебника, которые тебе дали: не добавляй фактов, которых там нет. Если "
    "фрагментов не хватает для темы — делай материалы проще и короче, но не выдумывай. "
    "В sources укажи номера страниц фрагментов, на которые опирался. Возвращай строго JSON."
)

SCHEMA = (
    '{"plan": {"title": "…", "goals": ["…"], "stages": [{"name": "…", "minutes": 5, "activity": "…"}], '
    '"homework": "…"}, "stories": [{"title": "…", "text": "…", "illustration": "…", '
    '"question": "…", "options": ["…", "…", "…"], "correct": 0}], '
    '"test": {"variants": [{"name": "A", "tasks": [{"text": "…", "points": 2, "answer": "…"}]}]}, '
    '"sources": [12, 13]}'
)


def build_prompt(topic: str, fragments: list[Fragment], variants: int, book: Textbook, lang: str) -> str:
    excerpt = "\n\n".join(f"[стр. {f.page}]\n{f.text}" for f in fragments)
    subject = SUBJECTS.get(book.subject or "", book.subject or "")
    return (
        f"Учебник: «{book.title}»{f', {subject}' if subject else ''}{f', {book.grade} класс' if book.grade else ''}.\n"
        f"Тема урока: «{topic}».\n\nФрагменты учебника:\n{excerpt}\n\n"
        "Подготовь:\n"
        f"1) plan — план урока на {LESSON_MINUTES} минут: 3–8 этапов (stages), сумма minutes = {LESSON_MINUTES}; "
        "goals — 2–4 цели; homework — домашнее задание;\n"
        "2) stories — Stories-урок из 8–12 коротких слайдов (text ≤ 3 предложений), illustration — что "
        "нарисовать на слайде; на 3–5 слайдах добавь question с 3 вариантами options и индексом correct;\n"
        f"3) test — контрольная ровно в {variants} вариантах одинаковой сложности, по 4–6 задач, "
        "у каждой points и краткий answer.\n"
        f"{language_rule(lang)}\n"
        f"Схема ответа: {SCHEMA}"
    )


class GeminiMaterialGenerator:
    def __init__(self, settings: Settings):
        from google import genai
        from google.genai import types

        self._types = types
        self.client = genai.Client(api_key=settings.gemini_api_key)
        self.model = settings.gemini_model

    async def generate(self, topic, fragments, variants, book, lang) -> dict:
        allowed = {f.page for f in fragments}
        last: Exception | None = None
        for _ in range(2):
            try:
                response = await self.client.aio.models.generate_content(
                    model=self.model,
                    contents=[build_prompt(topic, fragments, variants, book, lang)],
                    config=self._types.GenerateContentConfig(
                        system_instruction=SYSTEM_PROMPT, response_mime_type="application/json", temperature=0.4
                    ),
                )
                result = validate_material(extract_json(response.text or ""), variants, allowed)
                if result:
                    return result
                raise GeminiError("Материалы не прошли проверку схемы")
            except Exception as exc:
                log.warning("Генерация материалов: %s", exc)
                last = exc
        raise GeminiError(f"Не удалось сгенерировать материалы: {last}")


class StubMaterialGenerator:
    """Без модели: собирает материалы из найденных фрагментов — для демо и тестов."""

    async def generate(self, topic, fragments, variants, book, lang) -> dict:
        sentences = [s.strip() for f in fragments for s in f.text.replace("\n", " ").split(".") if len(s.strip()) > 20]
        sentences = (sentences or [topic]) * 12
        data = {
            "plan": {
                "title": topic,
                "goals": [f"Понять тему «{topic}»", "Научиться применять правило"],
                "stages": [
                    {"name": "Организационный момент", "minutes": 3, "activity": "Приветствие"},
                    {"name": "Объяснение", "minutes": 17, "activity": sentences[0]},
                    {"name": "Практика", "minutes": 20, "activity": "Решение задач"},
                    {"name": "Итог", "minutes": 5, "activity": "Рефлексия"},
                ],
                "homework": "Повторить параграф",
            },
            "stories": [
                {"title": f"Слайд {i + 1}", "text": sentences[i], "illustration": "схема",
                 **({"question": "Верно ли утверждение?", "options": ["Да", "Нет"], "correct": 0} if i % 3 == 0 else {})}
                for i in range(9)
            ],
            "test": {"variants": [
                {"name": chr(ord("A") + v), "tasks": [{"text": f"Задача {n + 1} по теме «{topic}»", "points": 2, "answer": str(n + v)} for n in range(4)]}
                for v in range(variants)
            ]},
            "sources": sorted({f.page for f in fragments}),
        }
        return validate_material(data, variants, {f.page for f in fragments})


def make_generator(settings: Settings | None = None) -> MaterialGenerator:
    settings = settings or get_settings()
    if settings.gemini_fake or not settings.gemini_api_key:
        return StubMaterialGenerator()
    return GeminiMaterialGenerator(settings)


# ---------- Учебники ----------

def _require_teacher(user: User) -> None:
    if user.role != "teacher":
        raise MaterialError("not_teacher", 403)


async def list_textbooks(session: AsyncSession, teacher: User) -> list[Textbook]:
    _require_teacher(teacher)
    rows = await session.scalars(
        select(Textbook)
        .where(or_(Textbook.owner_teacher_id.is_(None), Textbook.owner_teacher_id == teacher.id))
        .order_by(Textbook.owner_teacher_id.is_(None).desc(), desc(Textbook.created_at))
    )
    # Из библиотеки платформы показываем только готовые учебники
    return [b for b in rows if b.owner_teacher_id is not None or b.status == "ready"]


async def textbook_for(session: AsyncSession, teacher: User, textbook_id: int) -> Textbook:
    _require_teacher(teacher)
    book = await session.get(Textbook, textbook_id)
    if book is None or book.owner_teacher_id not in (None, teacher.id):
        raise MaterialError("textbook_not_found", 404)
    return book


async def create_textbook(
    session: AsyncSession, teacher: User, *, title: str, subject: str | None, grade: int | None, license_note: str
) -> Textbook:
    _require_teacher(teacher)
    book = Textbook(
        owner_teacher_id=teacher.id,
        title=title.strip()[:255],
        subject=subject,
        grade=grade,
        license_note=license_note.strip()[:500],
        pages=0,
        status="processing",
    )
    session.add(book)
    await session.commit()
    return book


async def delete_textbook(session: AsyncSession, teacher: User, textbook_id: int) -> None:
    book = await textbook_for(session, teacher, textbook_id)
    if book.owner_teacher_id != teacher.id:
        raise MaterialError("forbidden", 403)  # библиотеку платформы учитель не удаляет
    await session.delete(book)
    await session.commit()


# ---------- Материалы ----------

async def generate_material(
    session: AsyncSession,
    teacher: User,
    *,
    textbook_id: int,
    topic: str,
    variants: int = 2,
    lang: str | None = None,
    generator: MaterialGenerator | None = None,
    embedder: Embedder | None = None,
    settings: Settings | None = None,
) -> LessonMaterial:
    settings = settings or get_settings()
    book = await textbook_for(session, teacher, textbook_id)
    topic = topic.strip()
    if not topic:
        raise MaterialError("need_topic", 422)
    if not 2 <= variants <= 4:
        raise MaterialError("bad_variants", 422)
    lang = lang or teacher.lang
    try:
        fragments = await retrieve(session, book, topic, embedder or make_embedder(settings), k=settings.rag_top_k)
    except RagError as exc:
        raise MaterialError(exc.code, exc.status) from exc
    if not fragments:
        raise MaterialError("topic_not_in_textbook", 422)  # в учебнике нет ничего по теме — не выдумываем
    try:
        content = await asyncio.wait_for(
            (generator or make_generator(settings)).generate(topic, fragments, variants, book, lang),
            timeout=settings.materials_timeout_seconds,
        )
    except asyncio.TimeoutError as exc:
        raise MaterialError("generation_timeout", 504) from exc
    except GeminiError as exc:
        raise MaterialError("generation_failed", 502) from exc
    material = LessonMaterial(
        teacher_user_id=teacher.id,
        textbook_id=book.id,
        topic=topic[:255],
        lang=lang,
        content={k: content[k] for k in ("plan", "stories", "test", "answer_key")},
        sources=content["sources"],
    )
    session.add(material)
    await session.commit()
    return material


async def material_for(session: AsyncSession, teacher: User, material_id: int) -> LessonMaterial:
    _require_teacher(teacher)
    material = await session.get(LessonMaterial, material_id)
    if material is None or material.teacher_user_id != teacher.id:
        raise MaterialError("material_not_found", 404)
    return material


async def list_materials(session: AsyncSession, teacher: User, limit: int = 50) -> list[LessonMaterial]:
    _require_teacher(teacher)
    rows = await session.scalars(
        select(LessonMaterial)
        .where(LessonMaterial.teacher_user_id == teacher.id)
        .order_by(desc(LessonMaterial.updated_at), desc(LessonMaterial.id))
        .limit(limit)
    )
    return list(rows)


async def update_material(session: AsyncSession, teacher: User, material_id: int, content: dict, topic: str | None = None) -> LessonMaterial:
    """Правка учителем: та же схема, ключ ответов пересобирается из задач."""
    material = await material_for(session, teacher, material_id)
    variants = len(((content or {}).get("test") or {}).get("variants") or [])
    clean = validate_material({**content, "sources": material.sources}, variants, None) if 2 <= variants <= 4 else None
    if clean is None:
        raise MaterialError("bad_material", 422)
    material.content = {k: clean[k] for k in ("plan", "stories", "test", "answer_key")}
    if topic and topic.strip():
        material.topic = topic.strip()[:255]
    await session.commit()
    return material


async def delete_material(session: AsyncSession, teacher: User, material_id: int) -> None:
    material = await material_for(session, teacher, material_id)
    await session.delete(material)
    await session.commit()


# ---------- Экспорт ----------

EXPORT_LABELS = {
    "ru": {"plan": "План урока", "goals": "Цели", "min": "мин", "homework": "Домашнее задание", "stories": "Stories-урок",
           "test": "Контрольная работа", "variant": "Вариант", "points": "балл.", "key": "Ключ ответов", "sources": "Источник: учебник, стр."},
    "uz": {"plan": "Dars rejasi", "goals": "Maqsadlar", "min": "daq", "homework": "Uyga vazifa", "stories": "Stories-dars",
           "test": "Nazorat ishi", "variant": "Variant", "points": "ball", "key": "Javoblar kaliti", "sources": "Manba: darslik, bet"},
    "en": {"plan": "Lesson plan", "goals": "Goals", "min": "min", "homework": "Homework", "stories": "Stories lesson",
           "test": "Test", "variant": "Variant", "points": "pts", "key": "Answer key", "sources": "Source: textbook, p."},
}


def export_docx(material: LessonMaterial, book_title: str | None) -> bytes:
    from docx import Document

    L = EXPORT_LABELS.get(material.lang, EXPORT_LABELS["ru"])
    c = material.content
    doc = Document()
    doc.add_heading(material.topic, 0)
    if book_title:
        doc.add_paragraph().add_run(f"{L['sources']} {', '.join(map(str, material.sources))} — «{book_title}»").italic = True

    doc.add_heading(L["plan"], 1)
    if c["plan"].get("goals"):
        doc.add_paragraph(L["goals"] + ":")
        for g in c["plan"]["goals"]:
            doc.add_paragraph(g, style="List Bullet")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("#", L["min"], "")):
        cell.text = text
    for i, st in enumerate(c["plan"]["stages"], start=1):
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = str(i), str(st["minutes"]), f"{st['name']}. {st['activity']}"
    if c["plan"].get("homework"):
        doc.add_paragraph(f"{L['homework']}: {c['plan']['homework']}")

    doc.add_heading(L["stories"], 1)
    for i, sl in enumerate(c["stories"], start=1):
        doc.add_heading(f"{i}. {sl['title']}", 3)
        doc.add_paragraph(sl["text"])
        if sl.get("question"):
            doc.add_paragraph(sl["question"])
            for j, o in enumerate(sl["options"]):
                doc.add_paragraph(f"{'ABCD'[j]}) {o}", style="List Bullet")

    for v in c["test"]["variants"]:
        doc.add_page_break()
        doc.add_heading(f"{L['test']} · {L['variant']} {v['name']}", 1)
        for i, t in enumerate(v["tasks"], start=1):
            doc.add_paragraph(f"{i}. {t['text']} ({t['points']} {L['points']})")

    doc.add_page_break()
    doc.add_heading(L["key"], 1)
    for row in c["answer_key"]:
        doc.add_paragraph(f"{L['variant']} {row['variant']}, №{row['n']}: {row['answer']}")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
